"""PDF/PPTX extraction, provenance, draft synthesis, and approval separation."""

import json
from pathlib import Path
import sys

import pytest


ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/company-profile-analyzer/scripts'))
sys.path.insert(0,str(ROOT/'skills/report-html-pptx/scripts'))

from source_registry import register_source  # noqa: E402
from document_extract import extract_document, render_preview  # noqa: E402
from design_analyzer import analyze_design, draft_design_profile  # noqa: E402
from writing_analyzer import analyze_writing  # noqa: E402
from build_company_pack import build_company_pack, change_review_status  # noqa: E402
from company_pack import load_company_pack  # noqa: E402
sys.path.insert(0,str(ROOT/'scripts'))
from export_company_plugin import export_company_plugin  # noqa: E402


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    doc=Presentation()
    slide=doc.slides.add_slide(doc.slide_layouts[6])
    textbox=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2))
    run=textbox.text_frame.paragraphs[0].add_run()
    run.text='A made-up manufacturer tests every product.'
    run.font.name='Arial'
    run.font.size=Pt(18)
    run.font.color.rgb=RGBColor(15,72,98)
    doc.save(path)


def run_analysis(tmp_path, visibility='private', allow_excerpt=False):
    source=tmp_path/'sample.pptx'
    make_pptx(source)
    registry=tmp_path/'registry.json'
    item=register_source(registry,source,title='Fictional sample',language='en-US',
                         document_family='executive_report',visibility=visibility,
                         rights='MIT' if allow_excerpt else 'unverified',
                         allow_excerpt_in_pack=allow_excerpt,
                         distribution_authorized=allow_excerpt)
    extracted=extract_document(source)
    design=analyze_design(extracted,source_id=item['source_id'],document_family='executive_report')
    writing=analyze_writing(extracted,source_id=item['source_id'],language='en-US',
                            document_family='executive_report',allow_excerpt=allow_excerpt)
    (tmp_path/'design.json').write_text(json.dumps(design))
    (tmp_path/'writing.json').write_text(json.dumps(writing))
    return item,registry,extracted,design,writing


def test_registered_pptx_geometry_and_provenance(tmp_path):
    item,registry,extracted,design,writing=run_analysis(tmp_path)
    assert item['source_id'].startswith('SRC-')
    assert extracted['pages'][0]['shapes'][0]['fonts']==['Arial']
    assert extracted['pages'][0]['shapes'][0]['bbox_pt'] is not None
    assert design['quantitative_status']=='observed_measurement'
    assert design['visual_interpretation']['status']=='UNVERIFIED'
    assert design['dominant_colors']
    assert writing['private_text_excluded'] is True
    assert writing['references']==[]
    with pytest.raises(ValueError):
        register_source(registry,tmp_path/'sample.pptx',title='Conflicting title',
                        language='en-US',document_family='executive_report',visibility='private')


def test_private_sources_stay_out_of_draft_pack(tmp_path):
    item,registry,_,_,_=run_analysis(tmp_path)
    target=tmp_path/'output'/'fictional-co'
    target.parent.mkdir()
    build_company_pack(registry_file=registry,design_analyses=[tmp_path/'design.json'],
                       writing_analyses=[tmp_path/'writing.json'],output=target,
                       pack_id='fictional-co',company_name='Fictional Co',version='1.0.0',
                       default_language='en-US',enabled_languages=['en-US'])
    pack=load_company_pack(target)
    assert pack.metadata['status']=='draft'
    assert pack.writing_references==()
    assert (target/'sources.json').is_file()
    redacted=(target/'sources.json').read_text()
    assert 'sample.pptx' not in redacted
    assert str(tmp_path) not in redacted
    assert item['source_id'] in redacted
    with pytest.raises(ValueError):
        change_review_status(target,reviewer='Reviewer',note='Approved',action='approve')
    assert change_review_status(target,reviewer='Reviewer',note='Layout inspected',action='review')['status']=='reviewed'
    assert change_review_status(target,reviewer='Owner',note='Approved for use',action='approve')['status']=='approved'


def test_authorized_public_excerpts_stay_source_linked(tmp_path):
    item,registry,_,_,writing=run_analysis(tmp_path,'public',True)
    assert writing['references'][0]['status']=='observed'
    target=tmp_path/'output'/'public-sample'
    target.parent.mkdir()
    build_company_pack(registry_file=registry,design_analyses=[tmp_path/'design.json'],
                       writing_analyses=[tmp_path/'writing.json'],output=target,
                       pack_id='public-sample',company_name='Public sample',version='1.0.0',
                       default_language='en-US',enabled_languages=['en-US'])
    assert load_company_pack(target).writing_references[0]['source_ids']==[item['source_id']]


def test_pdf_extraction_when_renderer_available(tmp_path):
    fitz=pytest.importorskip('fitz')
    pdf=tmp_path/'sample.pdf'
    doc=fitz.open()
    page=doc.new_page(width=600,height=400)
    page.insert_text((60,80),'Quarterly operational update',fontsize=18)
    doc.save(pdf)
    doc.close()
    extracted=extract_document(pdf)
    assert extracted['pages'][0]['size_pt']==[600.0,400.0]
    assert any('Quarterly' in s['text'] for s in extracted['pages'][0]['shapes'])
    preview=render_preview(pdf,tmp_path/'preview')
    assert len(preview)==1 and Path(preview[0]).is_file()


def test_authorized_source_to_installable_plugin_end_to_end(tmp_path):
    item,registry,_,_,_=run_analysis(tmp_path,'public',True)
    target=tmp_path/'new-company'
    build_company_pack(registry_file=registry,design_analyses=[tmp_path/'design.json'],
                       writing_analyses=[tmp_path/'writing.json'],output=target,
                       pack_id='new-company',company_name='New Company',version='1.0.0',
                       default_language='en-US',enabled_languages=['en-US'])
    change_review_status(target,reviewer='Analyst',note='Measured source checked',action='review')
    change_review_status(target,reviewer='Owner',note='Rights and profile approved',action='approve')
    zip_path=tmp_path/'new-company.zip'
    result=export_company_plugin(target,zip_path,language_skills='none')
    assert len(result['skills'])==8
    import zipfile
    with zipfile.ZipFile(zip_path) as archive:
        combined=''.join(archive.read(name).decode('utf8') for name in archive.namelist()
                         if name.endswith(('.json','.md','.py','.yaml')))
        assert str(tmp_path) not in combined
        assert 'sample.pptx' not in combined
        assert item['source_id'] in combined
